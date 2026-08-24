#!/usr/bin/env python3
"""Check a pss4 sweep every ten minutes and build its two-sheet XLSX."""

from __future__ import annotations

import argparse
import ctypes
import json
import os
from pathlib import Path
import time

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


OMEGA_COLUMN = "omega (rad/s)"
TRANSMISSION_COLUMN = "transmission"
SPECTRAL_COLUMN = "spectral_flux (W/m2/K/(rad/s))"
EXPECTED_COATINGS = tuple(range(46))
EXPECTED_FREQUENCIES = 2502


def timestamp() -> str:
    return time.strftime("%Y-%m-%d %H:%M:%S")


def probe_csv(result_dir: Path) -> Path | None:
    paths = sorted(result_dir.glob("tot_P*.csv"))
    return paths[0] if len(paths) == 1 else None


def completed_cases(run_root: Path) -> list[int]:
    complete = []
    for coat_nm in EXPECTED_COATINGS:
        result_dir = run_root / "results" / f"coat_{coat_nm:03d}nm"
        path = probe_csv(result_dir)
        if path is not None and (result_dir / "final_report.txt").exists():
            complete.append(coat_nm)
    return complete


def write_status(run_root: Path, status: str, complete: list[int], **extra) -> None:
    payload = {
        "checked_at": timestamp(),
        "status": status,
        "completed_cases": len(complete),
        "total_cases": len(EXPECTED_COATINGS),
        "completed_coating_nm": complete,
        **extra,
    }
    target = run_root / "monitor_status.json"
    temporary = target.with_suffix(".json.tmp")
    temporary.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    os.replace(temporary, target)
    with (run_root / "monitor_10min.log").open("a", encoding="utf-8") as handle:
        handle.write(
            f"{payload['checked_at']} status={status} "
            f"completed={len(complete)}/{len(EXPECTED_COATINGS)}"
            + (f" xlsx={extra['xlsx_path']}" if "xlsx_path" in extra else "")
            + "\n"
        )


def load_and_validate(run_root: Path) -> tuple[np.ndarray, dict[int, pd.DataFrame]]:
    frames: dict[int, pd.DataFrame] = {}
    reference_omega = None
    for coat_nm in EXPECTED_COATINGS:
        result_dir = run_root / "results" / f"coat_{coat_nm:03d}nm"
        path = probe_csv(result_dir)
        if path is None:
            raise ValueError(f"Missing unique tot_P*.csv in {result_dir}")
        frame = pd.read_csv(path).sort_values(OMEGA_COLUMN).reset_index(drop=True)
        required = {OMEGA_COLUMN, TRANSMISSION_COLUMN, SPECTRAL_COLUMN}
        missing = required.difference(frame.columns)
        if missing:
            raise ValueError(f"{path} is missing columns {sorted(missing)}")
        if len(frame) != EXPECTED_FREQUENCIES:
            raise ValueError(f"{path} has {len(frame)} frequencies, expected {EXPECTED_FREQUENCIES}")
        values = frame[[OMEGA_COLUMN, TRANSMISSION_COLUMN, SPECTRAL_COLUMN]].to_numpy(float)
        if not np.all(np.isfinite(values)):
            raise ValueError(f"Non-finite values found in {path}")
        omega = frame[OMEGA_COLUMN].to_numpy(float)
        if reference_omega is None:
            reference_omega = omega
        elif not np.array_equal(reference_omega, omega):
            raise ValueError(f"Frequency grid mismatch in {path}")
        frames[coat_nm] = frame
    return reference_omega, frames


def populate_sheet(sheet, omega: np.ndarray, frames: dict[int, pd.DataFrame],
                   value_column: str) -> None:
    headers = [OMEGA_COLUMN] + [f"t_c={coat_nm} nm" for coat_nm in EXPECTED_COATINGS]
    sheet.append(headers)
    for row_index, frequency in enumerate(omega):
        sheet.append([float(frequency)] + [
            float(frames[coat_nm].at[row_index, value_column])
            for coat_nm in EXPECTED_COATINGS
        ])
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(omega) + 1}"
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["A"].width = 22
    for column in range(2, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 15
    sheet.row_dimensions[1].height = 24
    scientific_format = "0.000000000000000E+00"
    for row in sheet.iter_rows(min_row=2, max_row=len(omega) + 1,
                               min_col=1, max_col=len(headers)):
        for cell in row:
            cell.number_format = scientific_format
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(omega) + 1}"


def build_and_verify_xlsx(run_root: Path, output_path: Path) -> None:
    omega, frames = load_and_validate(run_root)
    workbook = Workbook()
    transmission = workbook.active
    transmission.title = "Transmission"
    spectral = workbook.create_sheet("Spectral Flux")
    populate_sheet(transmission, omega, frames, TRANSMISSION_COLUMN)
    populate_sheet(spectral, omega, frames, SPECTRAL_COLUMN)
    workbook.calculation.fullCalcOnLoad = True
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_suffix(".xlsx.tmp")
    workbook.save(temporary)
    os.replace(temporary, output_path)

    check = load_workbook(output_path, read_only=True, data_only=False)
    if check.sheetnames != ["Transmission", "Spectral Flux"]:
        raise ValueError(f"Unexpected sheet order: {check.sheetnames}")
    for sheet_name, source_column in (
        ("Transmission", TRANSMISSION_COLUMN),
        ("Spectral Flux", SPECTRAL_COLUMN),
    ):
        sheet = check[sheet_name]
        if sheet.max_row != EXPECTED_FREQUENCIES + 1 or sheet.max_column != 47:
            raise ValueError(
                f"{sheet_name} shape is {sheet.max_row}x{sheet.max_column}, expected 2503x47")
        if sheet["A1"].value != OMEGA_COLUMN or sheet["AU1"].value != "t_c=45 nm":
            raise ValueError(f"Unexpected headers in {sheet_name}")
        if sheet["A2"].value != float(omega[0]):
            raise ValueError(f"First frequency mismatch in {sheet_name}")
        expected = float(frames[45].at[0, source_column])
        if sheet["AU2"].value != expected:
            raise ValueError(f"Endpoint value mismatch in {sheet_name}")
    check.close()


def process_is_running(pid: int | None) -> bool:
    if pid is None:
        return True
    if os.name == "nt":
        process_query_limited_information = 0x1000
        still_active = 259
        handle = ctypes.windll.kernel32.OpenProcess(
            process_query_limited_information, False, pid)
        if not handle:
            return False
        try:
            exit_code = ctypes.c_ulong()
            if not ctypes.windll.kernel32.GetExitCodeProcess(handle, ctypes.byref(exit_code)):
                return False
            return exit_code.value == still_active
        finally:
            ctypes.windll.kernel32.CloseHandle(handle)
    try:
        os.kill(pid, 0)
        return True
    except OSError:
        return False


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("run_root", type=Path)
    parser.add_argument("output_xlsx", type=Path)
    parser.add_argument("--poll-seconds", type=float, default=600.0)
    parser.add_argument("--production-pid", type=int)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    run_root = args.run_root.resolve()
    output_path = args.output_xlsx.resolve()
    while True:
        complete = completed_cases(run_root)
        if len(complete) == len(EXPECTED_COATINGS):
            write_status(run_root, "building_xlsx", complete)
            build_and_verify_xlsx(run_root, output_path)
            write_status(run_root, "complete", complete, xlsx_path=str(output_path))
            return 0
        if not process_is_running(args.production_pid):
            write_status(run_root, "production_stopped_incomplete", complete)
            return 2
        write_status(run_root, "running", complete)
        time.sleep(args.poll_seconds)


if __name__ == "__main__":
    raise SystemExit(main())
